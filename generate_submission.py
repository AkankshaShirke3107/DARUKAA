import docx
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

def set_style(run, font_name="Calibri", font_size=11, bold=False, color=None):
    run.font.name = font_name
    run.font.size = Pt(font_size)
    run.font.bold = bold
    if color:
        run.font.color.rgb = color

def create_submission_doc():
    doc = docx.Document()
    
    # Set default margins
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    primary_color = RGBColor(46, 125, 50) # Dark green
    text_color = RGBColor(51, 51, 51)
    
    # Title
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_title = p_title.add_run("Darukaa.Earth — Biodiversity Intelligence System")
    set_style(run_title, font_size=18, bold=True, color=primary_color)
    
    p_subtitle = doc.add_paragraph()
    p_subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_sub = p_subtitle.add_run("AI Biodiversity Intelligence Chatbot Challenge Submission")
    set_style(run_sub, font_size=12, color=text_color)
    
    doc.add_paragraph() # Spacer

    # 1. Project Information
    h1 = doc.add_heading("1. Project Information", level=1)
    for run in h1.runs: set_style(run, font_size=14, bold=True, color=primary_color)
    
    desc = doc.add_paragraph()
    desc.add_run("Darukaa.Earth is a deterministic environmental intelligence system that combines a structured environmental knowledge layer, evidence retrieval, multi-variable ecological reasoning, conversational context, and evidence-backed biodiversity recommendations. It relies on a traceable reasoning engine rather than pure LLM generation.")
    
    info_table = doc.add_table(rows=6, cols=2)
    info_table.style = 'Light Shading Accent 3'
    info_data = [
        ("Applicant Name", "Akanksha Shirke"),
        ("Email", "akankshashirke3107@gmail.com"),
        ("GitHub Repository", "https://github.com/AkankshaShirke3107/DARUKAA"),
        ("Repository Visibility", "Public"),
        ("Live Demo URL", "N/A"),
        ("API / Docs URL", "N/A (Local /docs via FastAPI)"),
    ]
    for i, (key, value) in enumerate(info_data):
        row = info_table.rows[i]
        row.cells[0].text = key
        row.cells[1].text = value
        row.cells[0].paragraphs[0].runs[0].font.bold = True
    
    doc.add_paragraph() # Spacer

    # 2. README Overview
    h2 = doc.add_heading("2. README Overview", level=1)
    for run in h2.runs: set_style(run, font_size=14, bold=True, color=primary_color)
    
    h2_arch = doc.add_heading("Architecture", level=2)
    for run in h2_arch.runs: set_style(run, font_size=12, bold=True)
    p_arch = doc.add_paragraph("Next.js Frontend → FastAPI Backend → Knowledge / Evidence Layer → Ecological Reasoning Engine → Evidence-Backed Recommendations")
    
    doc.add_paragraph("The application separates the UI from the intelligence layer. The backend uses Pydantic for validation, a custom Knowledge Store for retrieving scientific evidence, a deterministic Interaction Detector for reasoning across variables, and a Recommendation Engine that templates evidence-backed interventions.")

    h2_db = doc.add_heading("Database / Knowledge Schema", level=2)
    for run in h2_db.runs: set_style(run, font_size=12, bold=True)
    doc.add_paragraph("The system currently utilizes an in-memory knowledge store and assessment cache, avoiding complex dependencies like PostgreSQL or pgvector. The knowledge base is a localized JSON structure containing 20 curated scientific documents from sources like FAO and IPCC. Evidence is retrieved via variable-overlap and keyword density scoring rather than vector embeddings. Assessment history and conversational state are maintained in memory during the runtime lifecycle.")

    h2_stack = doc.add_heading("Technology Stack", level=2)
    for run in h2_stack.runs: set_style(run, font_size=12, bold=True)
    doc.add_paragraph("Frontend: Next.js 16.3, React 19, Tailwind CSS, Recharts, Lucide React\nBackend: FastAPI, Uvicorn, Pydantic, google-generativeai, Pytest")

    h2_setup = doc.add_heading("Local Setup", level=2)
    for run in h2_setup.runs: set_style(run, font_size=12, bold=True)
    p_setup = doc.add_paragraph()
    p_setup.add_run("Frontend Setup:\n").bold = True
    p_setup.add_run("npm install\nnpm run dev\n")
    p_setup.add_run("Backend Setup:\n").bold = True
    p_setup.add_run("cd backend\ncp .env.example .env\npython -m venv venv\nsource venv/Scripts/activate\npip install -r requirements.txt\npython -m uvicorn app.main:app --host 0.0.0.0 --port 8000 --reload")

    p_cicd = doc.add_paragraph()
    p_cicd.add_run("CI/CD: ").bold = True
    p_cicd.add_run("Not currently configured.")
    
    doc.add_paragraph() # Spacer

    # 3. Review / Run Notes
    h3 = doc.add_heading("3. Review / Run Notes", level=1)
    for run in h3.runs: set_style(run, font_size=14, bold=True, color=primary_color)
    notes = [
        "Frontend URL: http://localhost:3000",
        "Backend / API URL: http://localhost:8000",
        "API Documentation: http://localhost:8000/docs",
        "Tests: Run 'python -m pytest tests/test_reasoning.py' from the backend directory",
        "Build: Run 'npm run build' from the root directory",
        "Environment Variables: Provide 'LLM_API_KEY' in the backend/.env file for natural language capabilities. Do NOT commit the .env file."
    ]
    for note in notes:
        doc.add_paragraph(note, style='List Bullet')
        
    doc.add_paragraph() # Spacer

    # 4. Challenge Alignment
    h4 = doc.add_heading("4. Challenge Alignment", level=1)
    for run in h4.runs: set_style(run, font_size=14, bold=True, color=primary_color)
    
    align_table = doc.add_table(rows=7, cols=2)
    align_table.style = 'Light Grid Accent 3'
    
    headers = align_table.rows[0].cells
    headers[0].text = "Requirement"
    headers[1].text = "Implementation"
    for cell in headers: cell.paragraphs[0].runs[0].font.bold = True
    
    align_data = [
        ("Knowledge System", "Actual implemented knowledge layer (20 scientific docs, keyword retrieval)"),
        ("Conversational Intelligence", "Implemented contextual conversation flow via backend /api/conversation/message"),
        ("Evidence-backed Recommendations", "Recommendations traceably linked to specific EVD- IDs and research papers"),
        ("Multi-Metric Reasoning", "Multi-variable ecological reasoning engine using 26-edge interaction graph"),
        ("Structured Input", "Implemented structured Pydantic assessment schema (soil, climate, land, biodiversity)"),
        ("Output Quality", "Actionable recommendations, impacted metrics, time horizons, and strict schemas")
    ]
    
    for i, (req, imp) in enumerate(align_data):
        row = align_table.rows[i+1]
        row.cells[0].text = req
        row.cells[1].text = imp

    # Save
    doc.save("Darukaa_Earth_Final_Submission.docx")

if __name__ == "__main__":
    create_submission_doc()
